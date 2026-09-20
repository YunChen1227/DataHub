#!/usr/bin/env python3
"""Query TSFX for every mobile in an xlsx (all sheets) and write responses back."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import time
from pathlib import Path

import requests
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

BASE_URL = os.environ.get("RELAY_BASE_URL", "http://aiszcloud.cn:8080")
DEFAULT_APP_KEY = "43fpitvnpvc6"
DEFAULT_APP_SECRET = "a7334ab983f34f255f88e092d5a74767"
MOBILE_RE = re.compile(r"^1\d{10}$")
MOBILE_HEADERS = {"手机号码", "手机号", "手机", "mobile", "电话", "电话号码"}

RESULT_HEADERS = (
    "策略poly",
    "errorCode",
    "bodyCode",
    "msg",
    "forbid",
    "命中说明",
    "callee",
    "logId",
    "uid",
    "reqid",
    "耗时ms",
    "原始响应",
)


def sign(body: dict[str, str], secret: str) -> str:
    items = sorted((k, v) for k, v in body.items() if v)
    return hashlib.md5(("".join(k + v for k, v in items) + secret).encode("utf-8")).hexdigest()


def forbid_text(value: object) -> str:
    if value == 1 or value == "1":
        return "命中"
    if value == 0 or value == "0":
        return "未命中"
    if value == -1 or value == "-1":
        return "无法判定"
    return ""


def find_mobile_col(header_row: tuple) -> int | None:
    for i, cell in enumerate(header_row):
        h = str(cell or "").strip().lower()
        if h in {x.lower() for x in MOBILE_HEADERS}:
            return i
    return None


def normalize_mobile(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        text = str(int(value))
    else:
        text = str(value).strip()
        if text.endswith(".0") and text[:-2].isdigit():
            text = text[:-2]
    return text


def query_one(mobile: str, poly: str, app_key: str, app_secret: str, session: requests.Session) -> dict:
    body = {"mobile": mobile, "poly": poly}
    payload = {
        "encryptionType": 1,
        "appKey": app_key,
        "sign": sign(body, app_secret),
        "body": body,
    }
    t0 = time.time()
    try:
        resp = session.post(f"{BASE_URL}/v1/openapi/zlx/querySrmxTSFX", json=payload, timeout=15)
        raw = resp.text
        data = resp.json()
    except Exception as exc:  # noqa: BLE001
        return {
            "error_code": "NETWORK",
            "body_code": "",
            "msg": str(exc),
            "forbid": "",
            "callee": "",
            "log_id": "",
            "uid": "",
            "reqid": "",
            "latency_ms": round((time.time() - t0) * 1000, 1),
            "raw": str(exc),
        }

    head = data.get("head") or {}
    body_obj = data.get("body") or {}
    result = body_obj.get("result") or {}
    forbid = ""
    callee = ""
    range_s = result.get("range") or ""
    try:
        parsed = json.loads(range_s) if range_s else []
        if isinstance(parsed, list) and parsed:
            forbid = parsed[0].get("forbid", "")
            callee = parsed[0].get("callee", "")
        elif isinstance(parsed, dict):
            forbid = parsed.get("forbid", "")
            callee = parsed.get("callee", "")
    except Exception:
        pass

    return {
        "error_code": str(head.get("errorCode", "")),
        "body_code": str(body_obj.get("code", "")),
        "msg": str(body_obj.get("msg") or head.get("errorMsg") or ""),
        "forbid": forbid,
        "callee": callee,
        "log_id": str(head.get("logId", "")),
        "uid": str(body_obj.get("uid", "")),
        "reqid": str(body_obj.get("reqid", "")),
        "latency_ms": head.get("time") if head.get("time") is not None else round((time.time() - t0) * 1000, 1),
        "raw": raw,
    }


def result_values(poly: str, r: dict) -> tuple:
    return (
        poly,
        r["error_code"],
        r["body_code"],
        r["msg"],
        r["forbid"],
        forbid_text(r["forbid"]),
        r["callee"],
        r["log_id"],
        r["uid"],
        r["reqid"],
        r["latency_ms"],
        r["raw"],
    )


def process_workbook(src: Path, poly: str, app_key: str, app_secret: str, out: Path | None = None) -> dict:
    wb = load_workbook(src)
    print(f"sheets ({len(wb.sheetnames)}): {wb.sheetnames}")

    summary = {
        "file": str(src),
        "sheets": len(wb.sheetnames),
        "sheet_names": list(wb.sheetnames),
        "poly": poly,
        "total": 0,
        "hit": 0,
        "miss": 0,
        "undetermined": 0,
        "error": 0,
        "by_sheet": {},
    }

    with requests.Session() as session:
        for sheet_name in wb.sheetnames:
            sh = wb[sheet_name]
            max_col = sh.max_column or 1
            first = next(sh.iter_rows(min_row=1, max_row=1, values_only=True), None) or ()
            mobile_col = find_mobile_col(first)
            start_row = 2 if mobile_col is not None else 1
            if mobile_col is None:
                mobile_col = 0

            result_start = max_col + 1
            for offset, title in enumerate(RESULT_HEADERS, start=0):
                sh.cell(1, result_start + offset, title)

            sheet_stats = {"rows": 0, "hit": 0, "miss": 0, "undetermined": 0, "error": 0}
            for row_idx in range(start_row, (sh.max_row or 1) + 1):
                mobile = normalize_mobile(sh.cell(row_idx, mobile_col + 1).value)
                if not mobile or mobile.lower() in {x.lower() for x in MOBILE_HEADERS}:
                    continue
                if not MOBILE_RE.match(mobile):
                    r = {
                        "error_code": "SKIP",
                        "body_code": "",
                        "msg": "非法手机号",
                        "forbid": "",
                        "callee": "",
                        "log_id": "",
                        "uid": "",
                        "reqid": "",
                        "latency_ms": "",
                        "raw": "",
                    }
                    print(f"[{sheet_name}#{row_idx}] skip invalid mobile={mobile}")
                else:
                    r = query_one(mobile, poly, app_key, app_secret, session)
                    print(
                        f"[{sheet_name}#{row_idx}] {mobile} -> ec={r['error_code']} "
                        f"bc={r['body_code']} forbid={r['forbid']} {forbid_text(r['forbid'])}"
                    )
                    time.sleep(0.15)

                values = result_values(poly, r)
                for offset, value in enumerate(values):
                    sh.cell(row_idx, result_start + offset, value)

                sheet_stats["rows"] += 1
                summary["total"] += 1
                if r["error_code"] != "0" or r["body_code"] != "001":
                    sheet_stats["error"] += 1
                    summary["error"] += 1
                elif r["forbid"] in (1, "1"):
                    sheet_stats["hit"] += 1
                    summary["hit"] += 1
                elif r["forbid"] in (0, "0"):
                    sheet_stats["miss"] += 1
                    summary["miss"] += 1
                elif r["forbid"] in (-1, "-1"):
                    sheet_stats["undetermined"] += 1
                    summary["undetermined"] += 1
                else:
                    sheet_stats["error"] += 1
                    summary["error"] += 1

            summary["by_sheet"][sheet_name] = sheet_stats

    dest = out or src
    try:
        wb.save(dest)
    except PermissionError:
        fallback = dest.with_name(f"{dest.stem}_tsfx结果{dest.suffix}")
        if fallback == dest:
            fallback = dest.with_name(f"{dest.stem}_out{dest.suffix}")
        wb.save(fallback)
        dest = fallback
        print(f"原文件被占用，已改写到: {dest}")
    summary["saved"] = str(dest)
    return summary


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("src", type=Path)
    parser.add_argument("--poly", default="C1", choices=["C1", "C2", "C3", "c1", "c2", "c3"])
    parser.add_argument("--app-key", default=os.environ.get("TSFX_APP_KEY", DEFAULT_APP_KEY))
    parser.add_argument("--app-secret", default=os.environ.get("TSFX_APP_SECRET", DEFAULT_APP_SECRET))
    parser.add_argument("--out", type=Path, default=None, help="output xlsx (default: overwrite src)")
    args = parser.parse_args()
    poly = args.poly.upper()

    health = requests.get(f"{BASE_URL}/healthz", timeout=10)
    print(f"healthz: {health.status_code} {health.text.strip()}")
    if health.status_code != 200:
        return 1

    print(f"target: {BASE_URL}/v1/openapi/zlx/querySrmxTSFX")
    print(f"appKey: {args.app_key}  poly: {poly}")
    print(f"file: {args.src}")

    summary = process_workbook(args.src, poly, args.app_key, args.app_secret, args.out)
    print("--- summary ---")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"saved: {summary.get('saved')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
