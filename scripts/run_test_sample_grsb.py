#!/usr/bin/env python3
"""Query GRSB route for an Excel sample and export 社保明细 xlsx."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import msoffcrypto
import os
import sys
import time
from pathlib import Path

import requests
from openpyxl import Workbook, load_workbook

sys.stdout.reconfigure(encoding="utf-8")

# grsb 专用凭证（勿与 grgjj 的 u4kwtvc588h5 / e1f1dd22… 混用）
DEFAULT_GRSB_APP_KEY = "bhiuvx5m4ug9"
DEFAULT_GRSB_APP_SECRET = "80474a2aa774f93d9cd0cd98a6643c08"

BASE_URL = os.environ.get("RELAY_BASE_URL", "http://aiszcloud.cn:8080")
ROOT = Path(__file__).resolve().parent.parent

HEADERS = ("姓名", "身份证号", "缴费时间", "参保状态", "缴费基数", "缴费单位", "个人身份")
CBJFZT = {"1": "正常参保", "2": "暂停参保"}
NAME_HEADERS = {"name", "姓名", "username"}
ID_HEADERS = {"idcard", "id_card", "身份证号", "证件号", "id号", "id"}


def sign(body: dict[str, str], secret: str) -> str:
    items = sorted((k, v) for k, v in body.items() if v)
    return hashlib.md5(("".join(k + v for k, v in items) + secret).encode()).hexdigest()


def open_sheet(path: Path, password: str):
    if password:
        with path.open("rb") as f:
            ms = msoffcrypto.OfficeFile(f)
            ms.load_key(password=password)
            buf = io.BytesIO()
            ms.decrypt(buf)
            buf.seek(0)
            return load_workbook(buf, read_only=True, data_only=True).active
    return load_workbook(path, read_only=True, data_only=True).active


def detect_columns(header: tuple) -> tuple[int, int] | None:
    """Return (name_col, id_col) from header row, or None if not a header."""
    cells = [str(c or "").strip().lower() for c in header]
    name_idx = id_idx = -1
    for i, h in enumerate(cells):
        if h in NAME_HEADERS:
            name_idx = i
        if h in ID_HEADERS:
            id_idx = i
    if name_idx >= 0 and id_idx >= 0:
        return name_idx, id_idx
    return None


def load_rows(path: Path, password: str) -> list[dict]:
    sh = open_sheet(path, password)
    rows: list[dict] = []
    name_col, id_col = 0, 1
    for i, row in enumerate(sh.iter_rows(values_only=True), 1):
        if not row:
            continue
        cols = detect_columns(row)
        if cols:
            name_col, id_col = cols
            continue
        if len(row) <= max(name_col, id_col):
            continue
        name = str(row[name_col]).strip() if row[name_col] is not None else ""
        id_card = str(row[id_col]).strip().upper() if row[id_col] is not None else ""
        if not name or not id_card:
            continue
        rows.append({"row": i, "name": name, "id_card": id_card})
    return rows


def query_one(rec: dict, session: requests.Session, app_key: str, app_secret: str) -> dict:
    body = {"name": rec["name"], "idCard": rec["id_card"]}
    payload = {
        "encryptionType": 1,
        "appKey": app_key,
        "sign": sign(body, app_secret),
        "body": body,
    }
    t0 = time.time()
    try:
        resp = session.post(
            f"{BASE_URL}/v1/openapi/zlx/querySrmxGRSB",
            json=payload,
            timeout=30,
        )
        raw = resp.text
        data = resp.json()
        head = data.get("head") or {}
        body_obj = data.get("body") or {}
        ec = str(head.get("errorCode", ""))
        bc = str(body_obj.get("code", ""))
        lat = float(head.get("time") or (time.time() - t0) * 1000)
        if ec == "0" and bc == "001":
            outcome = "found"
        elif ec == "0" and bc == "999":
            outcome = "not_found"
        elif ec == "0":
            outcome = "other_ok"
        else:
            outcome = "error"
        return {
            **rec,
            "http_status": resp.status_code,
            "error_code": ec,
            "body_code": bc,
            "latency_ms": lat,
            "outcome": outcome,
            "raw": raw,
        }
    except Exception as exc:  # noqa: BLE001
        return {
            **rec,
            "http_status": 0,
            "error_code": "NETWORK",
            "body_code": "",
            "latency_ms": None,
            "outcome": "error",
            "raw": str(exc),
        }


def parse_range(raw: str) -> dict | None:
    try:
        body = json.loads(raw).get("body") or {}
        rs = (body.get("result") or {}).get("range") or ""
        if not str(rs).strip():
            return None
        obj = json.loads(rs)
        return obj[0] if isinstance(obj, list) else obj
    except Exception:
        return None


def map_cbjfzt(value: object) -> str | None:
    if value is None or value == "":
        return None
    text = str(value).strip()
    return CBJFZT.get(text, text)


def result_row(rec: dict, raw: str | None) -> tuple:
    if not raw:
        return (rec["name"], rec["id_card"], None, None, None, None, None)
    data = parse_range(raw)
    if not data:
        return (rec["name"], rec["id_card"], None, None, None, None, None)
    return (
        data.get("xm") or rec["name"],
        data.get("sfz") or rec["id_card"],
        data.get("jfsj") or None,
        map_cbjfzt(data.get("cbjfzt")),
        data.get("jfjs") or None,
        data.get("jfdw") or None,
        data.get("grsf") or None,
    )


def main() -> int:
    parser = argparse.ArgumentParser(description="Batch query GRSB and export 社保明细")
    parser.add_argument("src", type=Path, help="source xlsx path")
    parser.add_argument("--password", default="", help="xlsx open password")
    parser.add_argument("--out-stem", help="output file stem (default: <src>_社保明细)")
    parser.add_argument("--app-key", default=os.environ.get("GRSB_APP_KEY", DEFAULT_GRSB_APP_KEY))
    parser.add_argument("--app-secret", default=os.environ.get("GRSB_APP_SECRET", DEFAULT_GRSB_APP_SECRET))
    parser.add_argument("--route", default="grsb", choices=["grsb"], help="sanity check: only grsb supported")
    args = parser.parse_args()

    src = args.src
    stem = args.out_stem or f"{src.stem}_社保明细"
    out_xlsx = ROOT / "test_res" / f"{stem}.xlsx"
    out_json = ROOT / "test_res" / f"{stem}_grsb_batch.json"
    app_key, app_secret = args.app_key, args.app_secret

    health = requests.get(f"{BASE_URL}/healthz", timeout=10)
    print(f"healthz: {health.status_code} {health.text.strip()}")
    if health.status_code != 200:
        return 1

    rows = load_rows(src, args.password)
    print(f"loaded {len(rows)} records from {src.name}")
    print(f"route: GRSB  target: {BASE_URL}")
    print(f"appKey(appid): {app_key}")
    print(f"appSecret: {app_secret[:6]}...{app_secret[-4:]}")

    results: list[dict] = []
    with requests.Session() as session:
        for i, rec in enumerate(rows, 1):
            r = query_one(rec, session, app_key, app_secret)
            results.append(r)
            print(
                f"[{i}/{len(rows)}] {rec['name']} {rec['id_card']} -> "
                f"{r['outcome']} ec={r['error_code']} bc={r['body_code']}"
            )
            time.sleep(0.2)

    results.sort(key=lambda x: x["row"])
    total = len(results)
    found = sum(1 for r in results if r["outcome"] == "found")
    not_found = sum(1 for r in results if r["outcome"] == "not_found")
    errors = total - found - not_found
    summary = {
        "route": "grsb",
        "app_key": app_key,
        "total": total,
        "found": found,
        "not_found": not_found,
        "error": errors,
        "found_rate": round(found / total * 100, 2) if total else 0,
        "not_found_rate": round(not_found / total * 100, 2) if total else 0,
        "error_rate": round(errors / total * 100, 2) if total else 0,
    }

    out_json.parent.mkdir(parents=True, exist_ok=True)
    out_json.write_text(
        json.dumps({"summary": summary, "results": results}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(HEADERS)
    with_data = 0
    for rec, r in zip(rows, results):
        out = result_row(rec, r["raw"])
        ws.append(out)
        if out[2] is not None:
            with_data += 1
    wb.save(out_xlsx)

    print("--- summary ---")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"with_data: {with_data}/{len(rows)}")
    print(f"saved json: {out_json}")
    print(f"saved xlsx: {out_xlsx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
